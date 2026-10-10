# Regenerates the derived acceptance TEST FIXTURES for G4 (AC-10, AC-11, AC-12, AC-14, AC-15).
#   python3 -I fixtures/acceptance/make_g4_fixtures.py     (needs reportlab + python-docx + LibreOffice `soffice`)
# Synthetic structural documents only — never a medical reference. The outputs are committed; tests never run this.
# Word-style documents (Arabic text, super/subscript FONT EFFECTS) are written as DOCX and converted with LibreOffice,
# the way an owner's own exam files are usually produced; exact layouts (options at the very bottom of a page) are
# drawn with reportlab.
import os
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = A4
FIX = 'TEST FIXTURE - synthetic structural document for automated tests. Not a medical reference.'
styles = getSampleStyleSheet()
P = lambda t: Paragraph(t, styles['Normal'])
HD = lambda t: Paragraph(t, styles['Heading2'])
ONLY = set(sys.argv[1:])


def wanted(name):
    return not ONLY or name in ONLY


def docx_to_pdf(doc, name):
    tmp = tempfile.mkdtemp()
    try:
        src = os.path.join(tmp, name.replace('.pdf', '.docx'))
        doc.save(src)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, src], check=True, capture_output=True, timeout=180)
        shutil.copy(os.path.join(tmp, name), os.path.join(HERE, name))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def rtl(p):
    """Right-to-left paragraph (bidi) with RTL runs, as Word writes Arabic."""
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement('w:bidi')
    b.set(qn('w:val'), '1')
    ppr.append(b)
    for r in p.runs:
        rpr = r._r.get_or_add_rPr()
        x = OxmlElement('w:rtl')
        x.set(qn('w:val'), '1')
        rpr.append(x)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ── AC-10: g4_long_questions.pdf — exact layout ────────────────────────────────────────────────────────────────────
# Q2's stem starts at the bottom of page 1 and stops mid-sentence («… his temperature is»); page 2 continues with a
# value at the line start («38.4 °C, …» — must not read as question 38) and the five options sit at the BOTTOM of
# page 2, after a large blank gap. Q3's options are split over pages 3 and 4. Running header + page numbers.
if wanted('g4_long_questions.pdf'):
    c = canvas.Canvas(os.path.join(HERE, 'g4_long_questions.pdf'), pagesize=A4)
    c.setTitle('G4 long questions (TEST FIXTURE)')
    state = {'y': 0}

    def page(n):
        c.setFont('Helvetica', 9)
        c.drawString(60, H - 40, 'G4 Long Question Bank - TEST FIXTURE')
        c.drawCentredString(W / 2, 30, str(n))
        state['y'] = H - 90

    def line(text, font='Helvetica', size=11, gap=16):
        c.setFont(font, size)
        c.drawString(60, state['y'], text)
        state['y'] -= gap

    page(1)
    line('G4 Long Question Bank', 'Helvetica-Bold', 15, 20)
    line(FIX, 'Helvetica', 9, 26)
    line('1. Which point is classically tender in acute appendicitis?')
    for o in ["A. Murphy's point", "B. McBurney's point", "C. Kehr's point", "D. Castell's point"]:
        line(o)
    state['y'] = 150
    line('2. A 58-year-old man is brought to the emergency department two hours after the onset of')
    line('severe epigastric pain radiating to the back. On examination his temperature is')
    c.showPage()
    page(2)
    line('38.4 °C, his blood pressure is 90/60 mmHg and his pulse is 118/min. Serum lipase is three')
    line('times the upper limit of normal. Which of the following is NOT part of the initial management?')
    state['y'] = 190
    for o in ['A. Intravenous crystalloid fluids', 'B. Adequate analgesia', 'C. Oxygen if hypoxic', 'D. Routine prophylactic antibiotics', 'E. Hourly urine output monitoring']:
        line(o)
    c.showPage()
    page(3)
    line('3. Which imaging test is first-line in suspected gallstones?')
    state['y'] = 120
    line('A. Plain abdominal X-ray')
    line('B. Ultrasound of the abdomen')
    c.showPage()
    page(4)
    line('C. MRCP')
    line('D. CT of the abdomen')
    state['y'] -= 20
    line('Answer Key', 'Helvetica-Bold', 13, 20)
    line('1. B 2. D 3. B')
    c.save()

# ── AC-10 + AC-11 (Arabic): g4_long_question_ar.pdf — Word document, RTL ─────────────────────────────────────────────
# An Arabic question whose stem crosses a page break and carries «11.5 ×10⁹/L» inside the Arabic sentence; its five
# Arabic-labelled options come after a large gap at the bottom of the next page(s). «عدا» = EXCEPT.
if wanted('g4_long_question_ar.pdf'):
    d = Document()
    d.add_paragraph('G4 Arabic Question Bank - TEST FIXTURE')
    d.add_paragraph(FIX)
    rtl(d.add_paragraph('1. أي مما يلي ليس من معايير ألفارادو؟'))
    for o in ['أ. هجرة الألم', 'ب. فقدان الشهية', 'ج. سكر الدم', 'د. ارتفاع الكريات البيض']:
        rtl(d.add_paragraph(o))
    for _ in range(14):
        d.add_paragraph('')
    rtl(d.add_paragraph('2. امرأة عمرها 30 سنة تراجع بألم في الحفرة الحرقفية اليمنى منذ 12 ساعة، وحرارتها 38.4 °C وتعداد الكريات البيض'))
    page_break(d)
    rtl(d.add_paragraph('11.5 ×10⁹/L. جميع ما يلي مناسب في التقييم الأولي عدا:'))
    for _ in range(22):
        d.add_paragraph('')
    for o in ['أ. اختبار الحمل', 'ب. تعداد الدم الكامل', 'ج. حقنة الباريوم الشرجية', 'د. فحص البول', 'هـ. الأمواج فوق الصوتية']:
        rtl(d.add_paragraph(o))
    page_break(d)
    rtl(d.add_paragraph('مفتاح الإجابة'))
    rtl(d.add_paragraph('1. ج 2. ج'))
    docx_to_pdf(d, 'g4_long_question_ar.pdf')

# ── AC-11: g4_units_negation.pdf — values typed with super/subscript FONT EFFECTS, decimal commas, bold NOT ──────────
if wanted('g4_units_negation.pdf'):
    doc = SimpleDocTemplate(os.path.join(HERE, 'g4_units_negation.pdf'), pagesize=A4, title='g4_units_negation')
    doc.build([
        HD('G4 Units Bank'), P(FIX), Spacer(1, 8),
        P('1. In suspected appendicitis, a white cell count of 11.5 &times; 10<super>9</super>/L:'),
        P('A. Confirms the diagnosis'), P('B. Supports but does not confirm the diagnosis'), P('C. Excludes the diagnosis'), P('D. Indicates perforation'),
        P('2. An arterial blood gas shows pH 7.32 and PaCO<sub>2</sub> 52 mmHg. Which disturbance is present?'),
        P('A. Respiratory acidosis'), P('B. Metabolic alkalosis'), P('C. Respiratory alkalosis'), P('D. Metabolic acidosis'),
        P('3. Serum potassium is 6,5 mmol/L and creatinine 1,2 mg/dL. Which is the first step?'),
        P('A. Calcium gluconate 10 mL of 10%'), P('B. Observation'), P('C. Oral potassium'), P('D. Repeat in 24 h'),
        P('4. Which of the following is <b>NOT</b> a feature of shock?'),
        P('A. Tachycardia'), P('B. Hypotension'), P('C. Warm peripheries in all cases'), P('D. Oliguria &lt; 0.5 mL/kg/h'),
        P('5. All of the following are risk factors for gallstones <u>except</u>:'),
        P('A. Female sex'), P('B. Obesity'), P('C. Rapid weight loss'), P('D. Regular physical activity'),
        Spacer(1, 12), HD('Answers'), P('1. B 2. A 3. A 4. C 5. D'),
    ])

# ── AC-11 (Arabic): g4_negation_ar.pdf — the negation words themselves are lam-alef ligatures («لا»، «إلا») ──────────
# LibreOffice's PDF text layer reverses every lam-alef: «لا» → «ال», «إلا» → «إال». Also «عدا», «خاطئة», an Arabic-Indic
# decimal «٣٫٥», and an English section typed in Word with super/subscript effects (10⁹, PaCO₂, HCO₃⁻).
if wanted('g4_negation_ar.pdf'):
    d = Document()
    d.add_paragraph('G4 Arabic negation bank - TEST FIXTURE')
    d.add_paragraph(FIX)
    qs = [
        ('1. أي مما يلي لا يسبب ارتفاع حرارة المريض؟', ['أ. التهاب الزائدة', 'ب. الخراج', 'ج. الراحة في السرير', 'د. التهاب الصفاق']),
        ('2. جميع ما يلي من مضاعفات التهاب الزائدة إلا:', ['أ. الانثقاب', 'ب. الخراج', 'ج. التهاب الصفاق', 'د. حصى المرارة']),
        ('3. كل ما يلي صحيح ما عدا:', ['أ. الألم يبدأ حول السرة', 'ب. يهاجر الألم إلى اليمين', 'ج. القيء يسبق الألم دائما', 'د. الحمى خفيفة عادة']),
        ('4. مريض صوديوم المصل لديه 128 ملمول/لتر والبوتاسيوم ٣٫٥ ملمول/لتر. أي العبارات التالية خاطئة؟', ['أ. الصوديوم منخفض', 'ب. البوتاسيوم طبيعي', 'ج. الصوديوم مرتفع', 'د. يلزم تقييم السوائل']),
    ]
    for stem, opts in qs:
        rtl(d.add_paragraph(stem))
        for o in opts:
            rtl(d.add_paragraph(o))
    rtl(d.add_paragraph('مفتاح الإجابة'))
    rtl(d.add_paragraph('1. ج 2. د 3. ج 4. ج'))
    d.add_paragraph('Section E')
    p = d.add_paragraph('5. A white cell count of 11.5 × 10')
    r = p.add_run('9')
    r.font.superscript = True
    p.add_run('/L in suspected appendicitis:')
    for o in ['A. Confirms the diagnosis', 'B. Supports but does not confirm it', 'C. Excludes it', 'D. Indicates perforation']:
        d.add_paragraph(o)
    p = d.add_paragraph('6. pH 7.32 with PaCO')
    r = p.add_run('2')
    r.font.subscript = True
    p.add_run(' 52 mmHg and HCO')
    r = p.add_run('3')
    r.font.subscript = True
    r = p.add_run('−')
    r.font.superscript = True
    p.add_run(' 26 mmol/L indicates:')
    for o in ['A. Respiratory acidosis', 'B. Metabolic alkalosis', 'C. Respiratory alkalosis', 'D. Metabolic acidosis']:
        d.add_paragraph(o)
    docx_to_pdf(d, 'g4_negation_ar.pdf')


# ── AC-12 ───────────────────────────────────────────────────────────────────────────────────────────────────────────
def q(n, stem, opts):
    return [P(f'{n}. {stem}')] + [P(f'{chr(65 + i)}. {o}') for i, o in enumerate(opts)]


SA = q(1, 'Which vitamin deficiency causes scurvy?', ['Vitamin A', 'Vitamin C', 'Vitamin D', 'Vitamin K']) + q(2, 'Which vitamin is fat soluble?', ['Vitamin B1', 'Vitamin B12', 'Vitamin C', 'Vitamin E'])
SB = (
    q(1, 'Which organ produces insulin?', ['Liver', 'Spleen', 'Pancreas', 'Kidney'])
    + q(2, 'Which hormone raises blood glucose?', ['Insulin', 'Glucagon', 'Somatostatin', 'Leptin'])
    + q(3, 'Which cell type secretes glucagon?', ['Alpha cells', 'Beta cells', 'Delta cells', 'PP cells'])
)
SC = q(1, 'Which bone is the longest in the body?', ['Femur', 'Tibia', 'Humerus', 'Radius'])

# g4_sections_merged_key.pdf: sections A, B, C all numbered from 1; the key is printed on a later page as two short
# lines (Section B FIRST) that the layout merges into ONE region «Section B: … Section A: …»; the Section A line
# keys a question 4 that does not exist; Section C has no key at all.
if wanted('g4_sections_merged_key.pdf'):
    doc = SimpleDocTemplate(os.path.join(HERE, 'g4_sections_merged_key.pdf'), pagesize=A4, title='g4_sections_merged_key')
    doc.build([HD('Sections Bank'), P(FIX), Spacer(1, 6), HD('Section A'), *SA, HD('Section B'), *SB, HD('Section C'), *SC, PageBreak(),
               HD('Answer Key'), P('Section B: 1. C 2. B 3. A'), P('Section A: 1. B 2. D 4. A')])

# g4_key_formats.pdf: Part 1 keyed «Q1: B Q2: D» (a layout that used to become a bogus question «B Q2: D»); Part 2's
# key printed in a layout the parser cannot read («Q1 is C, Q2 is B, Q3 is A») → reported, never guessed.
if wanted('g4_key_formats.pdf'):
    SA2 = q(1, 'Which vitamin deficiency causes rickets in children?', ['Vitamin A', 'Vitamin D', 'Vitamin C', 'Vitamin K']) + q(2, 'Which vitamin is needed for clotting factor synthesis?', ['Vitamin B1', 'Vitamin C', 'Vitamin E', 'Vitamin K'])
    SB2 = q(1, 'Which gland secretes thyroxine?', ['Adrenal', 'Pituitary', 'Thyroid', 'Pineal']) + q(2, 'Which hormone lowers blood glucose?', ['Glucagon', 'Insulin', 'Cortisol', 'Adrenaline']) + q(3, 'Which organ stores glycogen mainly?', ['Liver', 'Spleen', 'Kidney', 'Lung'])
    doc = SimpleDocTemplate(os.path.join(HERE, 'g4_key_formats.pdf'), pagesize=A4, title='g4_key_formats')
    doc.build([HD('Key Formats Bank'), P(FIX), Spacer(1, 6), HD('Part 1'), *SA2, HD('Answer key'), P('Q1: B Q2: D'), HD('Part 2'), *SB2, HD('Answer key'), P('Q1 is C, Q2 is B, Q3 is A')])

# g4_sections_inline_keys.pdf: «Answers» printed right after EACH part (Part 1, then Part 2 at the end of the file).
# Part 1's key sits between the two parts (bound to Part 1); Part 2's trailing, unlabeled key could also be a key for
# the whole file → it is kept unbound with a review item (never applied by number alone).
if wanted('g4_sections_inline_keys.pdf'):
    PA = q(1, 'Which nerve supplies the diaphragm?', ['Vagus', 'Phrenic', 'Accessory', 'Hypoglossal']) + q(2, 'Which artery supplies the left ventricle mainly?', ['Right coronary', 'Left anterior descending', 'Internal thoracic', 'Subclavian'])
    PB = q(1, 'Which bone forms the heel?', ['Talus', 'Navicular', 'Calcaneus', 'Cuboid']) + q(2, 'Which muscle flexes the elbow mainly?', ['Triceps', 'Brachialis', 'Deltoid', 'Supinator']) + q(3, 'Which joint is a ball and socket joint?', ['Hip', 'Knee', 'Elbow', 'Ankle'])
    doc = SimpleDocTemplate(os.path.join(HERE, 'g4_sections_inline_keys.pdf'), pagesize=A4, title='g4_sections_inline_keys')
    doc.build([HD('Inline Keys Bank'), P(FIX), Spacer(1, 6), HD('Part 1'), *PA, HD('Answers'), P('1. B 2. B'), HD('Part 2'), *PB, HD('Answers'), P('1. C 2. B 3. A')])

# g4_sections_ar.pdf: Arabic sections «القسم الأول» / «القسم الثاني» both numbered from 1; the key lists the SECOND
# section first.
if wanted('g4_sections_ar.pdf'):
    d = Document()
    d.add_paragraph('G4 Arabic sections bank - TEST FIXTURE')
    d.add_paragraph(FIX)
    rtl(d.add_paragraph('القسم الأول'))
    for s in ['1. ما الفيتامين الذي يسبب نقصه داء الإسقربوط؟', 'أ. فيتامين A', 'ب. فيتامين C', 'ج. فيتامين D', 'د. فيتامين K',
              '2. أي الفيتامينات التالية يذوب في الدهون؟', 'أ. فيتامين B1', 'ب. فيتامين B12', 'ج. فيتامين C', 'د. فيتامين E']:
        rtl(d.add_paragraph(s))
    rtl(d.add_paragraph('القسم الثاني'))
    for s in ['1. أي عضو يفرز الإنسولين؟', 'أ. الكبد', 'ب. الطحال', 'ج. البنكرياس', 'د. الكلية',
              '2. أي هرمون يرفع سكر الدم؟', 'أ. الإنسولين', 'ب. الغلوكاغون', 'ج. السوماتوستاتين', 'د. اللبتين']:
        rtl(d.add_paragraph(s))
    rtl(d.add_paragraph('مفتاح الإجابة'))
    rtl(d.add_paragraph('القسم الثاني: 1. ج 2. ب'))
    rtl(d.add_paragraph('القسم الأول: 1. ب 2. د'))
    docx_to_pdf(d, 'g4_sections_ar.pdf')

# ── AC-15 ───────────────────────────────────────────────────────────────────────────────────────────────────────────
# g4_wrong_key.pdf: a printed key that the course lecture (lecture_appendicitis.pdf: pain «migrates to the right iliac
# fossa (McBurney's point)») contradicts — key A (Murphy's point). Q2 is keyed in agreement with the lecture.
if wanted('g4_wrong_key.pdf'):
    doc = SimpleDocTemplate(os.path.join(HERE, 'g4_wrong_key.pdf'), pagesize=A4, title='g4_wrong_key')
    doc.build([
        HD('Key Check Bank'), P(FIX), Spacer(1, 8),
        P('1. In acute appendicitis, to which point does the pain classically migrate?'),
        P("A. Murphy's point"), P("B. McBurney's point"), P("C. Kehr's point"), P("D. Castell's point"),
        P('2. Which test is required in a woman of reproductive age with suspected appendicitis?'),
        P('A. Serum amylase'), P('B. Pregnancy test (beta-hCG)'), P('C. Barium enema'), P('D. Colonoscopy'),
        Spacer(1, 12), HD('Answer Key'), P('1. A 2. B'),
    ])

# g4_a1_other_key.pdf: the golden question A1 printed again, word for word, with ANOTHER key (C) — a second source
# that disagrees AFTER the owner already answered A1 (conflict arriving later, AC-15 + AC-26).
if wanted('g4_a1_other_key.pdf'):
    doc = SimpleDocTemplate(os.path.join(HERE, 'g4_a1_other_key.pdf'), pagesize=A4, title='g4_a1_other_key')
    doc.build([
        HD('Revision Sheet'), P(FIX), Spacer(1, 8),
        P('1. Which point is classically tender in acute appendicitis?'),
        P("A. Murphy's point"), P("B. McBurney's point"), P("C. Kehr's point"), P("D. Castell's point"),
        Spacer(1, 12), HD('Answer Key'), P('1. C'),
    ])

print('ok')
