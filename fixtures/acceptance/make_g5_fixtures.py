# Regenerates the derived acceptance TEST FIXTURES for G5 (AC-16, AC-17, AC-19).
#   python3 -I fixtures/acceptance/make_g5_fixtures.py     (needs reportlab + Pillow + python-docx + LibreOffice `soffice`)
# Synthetic structural documents only — never a medical reference. The outputs are committed; tests never run this.
# Symbols that carry meaning (U+2212 minus, ♀ / ♂, ✓) are drawn with DejaVu Sans (embedded, with a ToUnicode map) so the
# PDF text layer holds the real characters; the Arabic documents are written as DOCX and converted with LibreOffice,
# the way an owner's own files are usually produced.
import io
import os
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = A4
FIX = 'TEST FIXTURE - synthetic structural document for automated tests. Not a medical reference.'
ONLY = set(sys.argv[1:])
pdfmetrics.registerFont(TTFont('DejaVu', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DejaVu-Bold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))


def wanted(name):
    return not ONLY or name in ONLY


class Sheet:
    """A4 pages drawn line by line with DejaVu Sans, a running header and a page number."""

    def __init__(self, name, header):
        self.c = canvas.Canvas(os.path.join(HERE, name), pagesize=A4)
        self.c.setTitle(f'{header} (TEST FIXTURE)')
        self.header = header
        self.n = 0
        self.y = 0
        self.new_page()

    def new_page(self):
        if self.n:
            self.c.showPage()
        self.n += 1
        self.c.setFont('DejaVu', 9)
        self.c.drawString(60, H - 40, f'{self.header} - TEST FIXTURE')
        self.c.drawCentredString(W / 2, 30, str(self.n))
        self.y = H - 90

    def line(self, text, font='DejaVu', size=11, gap=17):
        self.c.setFont(font, size)
        self.c.drawString(60, self.y, text)
        self.y -= gap

    def image(self, png_bytes, w, h):
        self.c.drawImage(ImageReader(io.BytesIO(png_bytes)), 60, self.y - h, width=w, height=h)
        self.y -= h + 14

    def save(self):
        self.c.save()


def synthetic_picture():
    """A drawn grey «scan» with an oval — labelled as a synthetic test drawing, no medical content."""
    img = Image.new('RGB', (640, 360), (36, 36, 40))
    d = ImageDraw.Draw(img)
    d.ellipse((180, 90, 460, 270), outline=(210, 210, 210), width=10, fill=(90, 90, 96))
    d.rectangle((20, 20, 620, 340), outline=(160, 160, 160), width=4)
    d.text((32, 30), 'TEST FIXTURE - synthetic drawing', fill=(230, 230, 230))
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return buf.getvalue()


# ── AC-19: g5_marked_bank.pdf — answer marks printed next to options, an inline key glued to the last option, and a
# picture question whose caption names the answer. The printed key at the end binds Q1, Q2 and Q4.
if wanted('g5_marked_bank.pdf'):
    s = Sheet('g5_marked_bank.pdf', 'G5 Marked Question Bank')
    s.line('G5 Marked Question Bank', 'DejaVu-Bold', 15, 20)
    s.line(FIX, 'DejaVu', 9, 26)
    s.line('1. Where is the tenderness classically located in acute appendicitis?')
    for o in ["A. Murphy's point", "B. McBurney's point ✓", "C. Kehr's point", "D. Castell's point"]:
        s.line(o)
    s.y -= 8
    s.line('2. Which imaging test is first-line in children with suspected appendicitis?')
    for o in ['A. Ultrasound *', 'B. CT abdomen', 'C. MRI pelvis', 'D. Plain abdominal X-ray']:
        s.line(o)
    s.y -= 8
    s.line('3. Which clinical sign is associated with acute cholecystitis?')
    for o in ["A. Murphy's sign", 'B. Psoas sign', 'C. Obturator sign', "D. Rovsing's sign      Answer: A"]:
        s.line(o)
    s.new_page()
    s.line('4. What does the image below show?')
    s.image(synthetic_picture(), 320, 180)
    s.line('Figure 1: Ultrasound of an inflamed appendix (synthetic TEST FIXTURE drawing).', 'DejaVu', 9, 22)
    for o in ['A. Inflamed appendix', 'B. Gallstones', 'C. Kidney stone', 'D. Normal bowel loop']:
        s.line(o)
    s.y -= 12
    s.line('Answers', 'DejaVu-Bold', 13, 20)
    s.line('1. B 2. A 4. A')
    s.save()

# ── AC-17: g5_signs_a.pdf / g5_signs_b.pdf — the same questions in two files, except for ONE symbol that changes the
# medical meaning (base excess «−8» vs «8» with the real U+2212 minus; «♀» vs «♂»). Each file keys its own answer.
# Q3 is the Golden Set's A1, word for word, in both files: the true exact duplicate (control).
SIGN_Q = [
    ('1. A blood gas shows a base excess of {be} mmol/L. Which disturbance does it indicate?',
     ['A. Metabolic acidosis', 'B. Metabolic alkalosis', 'C. Respiratory acidosis', 'D. Respiratory alkalosis']),
    ('2. A 24-year-old {sex} presents with right iliac fossa pain. Which test should be done first?',
     ['A. Pregnancy test (β-hCG)', 'B. Urinalysis', 'C. Barium enema', 'D. Colonoscopy']),
    ('3. Which point is classically tender in acute appendicitis?',
     ["A. Murphy's point", "B. McBurney's point", "C. Kehr's point", "D. Castell's point"]),
]
for name, be, sex, keys in [('g5_signs_a.pdf', '−8', '♀', '1. A 2. A 3. B'), ('g5_signs_b.pdf', '8', '♂', '1. B 2. B 3. B')]:
    if not wanted(name):
        continue
    s = Sheet(name, 'G5 Signs Bank ' + name[-5].upper())
    s.line('G5 Signs Bank', 'DejaVu-Bold', 15, 20)
    s.line(FIX, 'DejaVu', 9, 26)
    for stem, opts in SIGN_Q:
        s.line(stem.format(be=be, sex=sex))
        for o in opts:
            s.line(o)
        s.y -= 8
    s.line('Answers', 'DejaVu-Bold', 13, 20)
    s.line(keys)
    s.save()


# ── AC-16 (Arabic): g5_questions_ar.pdf first, g5_lecture_ar.pdf later — Word documents, RTL ─────────────────────────
def docx_to_pdf(doc, name):
    tmp = tempfile.mkdtemp()
    try:
        src = os.path.join(tmp, name.replace('.pdf', '.docx'))
        doc.save(src)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, src], check=True, capture_output=True, timeout=180)
        shutil.copy(os.path.join(tmp, name), os.path.join(HERE, name))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def rtl(p):
    ppr = p._p.get_or_add_pPr()
    b = OxmlElement('w:bidi')
    b.set(qn('w:val'), '1')
    ppr.append(b)
    for r in p.runs:
        rpr = r._r.get_or_add_rPr()
        x = OxmlElement('w:rtl')
        x.set(qn('w:val'), '1')
        rpr.append(x)
    return p


if wanted('g5_questions_ar.pdf'):
    d = Document()
    d.add_paragraph('G5 Arabic Question Bank - TEST FIXTURE')
    d.add_paragraph(FIX)
    qs = [
        ('1. ما العلامة السريرية المرتبطة بالتهاب المرارة الحاد؟', ['أ. علامة مورفي', 'ب. علامة روفسينغ', 'ج. علامة بسواس', 'د. علامة كير']),
        ('2. ما الفحص الأولي المفضل عند الشك بحصى المرارة؟', ['أ. Ultrasound', 'ب. CT abdomen', 'ج. MRCP', 'د. Plain X-ray']),
        ('3. ما العلاج الأولي لكسر عظم الفخذ المغلق عند البالغين؟', ['أ. التثبيت الجراحي', 'ب. الراحة فقط', 'ج. المضادات الحيوية', 'د. العلاج الطبيعي']),
    ]
    for stem, opts in qs:
        rtl(d.add_paragraph(stem))
        for o in opts:
            rtl(d.add_paragraph(o))
    rtl(d.add_paragraph('الإجابات'))
    rtl(d.add_paragraph('1. أ 2. أ 3. أ'))
    docx_to_pdf(d, 'g5_questions_ar.pdf')

if wanted('g5_lecture_ar.pdf'):
    d = Document()
    d.add_paragraph('G5 Arabic lecture - TEST FIXTURE')
    d.add_paragraph(FIX)
    rtl(d.add_heading('التهاب المرارة الحاد', level=1))
    rtl(d.add_paragraph('يبدأ ألم المرارة في الربع العلوي الأيمن من البطن وقد ينتشر إلى الكتف الأيمن.'))
    rtl(d.add_paragraph('علامة مورفي إيجابية عادة في التهاب المرارة الحاد، وتُفحص بالضغط تحت الحافة الضلعية اليمنى أثناء الشهيق.'))
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    rtl(d.add_heading('الفحوصات', level=1))
    rtl(d.add_paragraph('الفحص الأولي المفضل عند الشك بحصى المرارة هو Ultrasound لأنه متاح وسريع ولا يعرّض المريض للأشعة.'))
    rtl(d.add_paragraph('يُستخدم MRCP عند الشك بحصى القناة الصفراوية المشتركة.'))
    docx_to_pdf(d, 'g5_lecture_ar.pdf')

# ── AC-17 (E2E): g5_dup_a.pdf (question source) / g5_dup_b.pdf (previous exam) — the SAME two questions in two files,
# used only by G5 (the Golden Set's A1 is shared by every spec on an E2E server, and another spec gives it a conflicting
# key). Both are answered by the Golden Set lecture (pregnancy test p. 12, ultrasound p. 12).
DUP_Q = [
    ('1. Which patients with suspected appendicitis need a pregnancy test (β-hCG)?',
     ['A. Women of reproductive age', 'B. Children under five', 'C. Men over sixty', 'D. Every adult patient']),
    ('2. Which imaging test is first-line in pregnant women with suspected appendicitis?',
     ['A. CT abdomen', 'B. Ultrasound', 'C. Plain abdominal X-ray', 'D. MRI pelvis']),
]
for name, header in [('g5_dup_a.pdf', 'G5 Duplicate Bank A'), ('g5_dup_b.pdf', 'G5 Previous Exam B')]:
    if not wanted(name):
        continue
    s = Sheet(name, header)
    s.line(header, 'DejaVu-Bold', 15, 20)
    s.line(FIX, 'DejaVu', 9, 26)
    for stem, opts in DUP_Q:
        s.line(stem)
        for o in opts:
            s.line(o)
        s.y -= 8
    s.line('Answers', 'DejaVu-Bold', 13, 20)
    s.line('1. A 2. B')
    s.save()
