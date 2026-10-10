# Regenerates the derived acceptance TEST FIXTURES for G8 (AC-26, AC-27, security sweep).
#   python3 -I fixtures/acceptance/make_g8_fixtures.py [file …]   (needs reportlab + python-pptx + LibreOffice `soffice`)
# Synthetic structural documents only — never a medical reference. The outputs are committed; tests never run this.
#
#  * g8_key_bank.pdf        a question source with five G8-only questions answered by the Golden Set appendicitis lecture;
#                           Q3's PRINTED key is wrong (B «Ultrasound»; the lecture: CT is preferred in adults) so the owner
#                           has a real key to correct (AC-26). G8-only wording: no other spec shares these questions.
#  * g8_linked_image.pptx   one slide whose picture is only LINKED (r:link, TargetMode="External") to
#                           http://127.0.0.1:65001/g8-pptx-ssrf.png — the reader's fixed rendering is made by LibreOffice,
#                           which would fetch it (SSRF from an uploaded file). Tests rewrite the port to a local trap.
#  * g8_linked_image.doc    a Word 97 document whose picture is only LINKED to http://127.0.0.1:65001/g8-doc-ssrf.png
#                           (legacy .doc is converted by LibreOffice). Tests replace the 5-digit port in place (OLE2 has
#                           no checksum over the stream text).
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN = os.path.join(HERE, '..', 'golden')
W, H = A4
FIX = 'TEST FIXTURE - synthetic structural document for automated tests. Not a medical reference.'
ONLY = set(sys.argv[1:])
TRAP = 'http://127.0.0.1:65001'
pdfmetrics.registerFont(TTFont('DejaVu', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DejaVu-Bold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))

# LibreOffice must never fetch anything while these fixtures are BUILT either: every proxy points at a closed port.
NO_NET = {
    'PATH': os.environ.get('PATH', '/usr/bin:/bin'),
    'LANG': 'C.UTF-8',
    'http_proxy': 'http://127.0.0.1:1',
    'https_proxy': 'http://127.0.0.1:1',
    'HTTP_PROXY': 'http://127.0.0.1:1',
    'HTTPS_PROXY': 'http://127.0.0.1:1',
    'no_proxy': '',
    'NO_PROXY': '',
}


def wanted(name):
    return not ONLY or name in ONLY


# ── g8_key_bank.pdf ──
BANK = [
    ('1. G8 set: where does the pain of acute appendicitis usually begin?',
     ['A. In the periumbilical region', 'B. In the left iliac fossa', 'C. In the epigastrium only', 'D. In the right loin']),
    ('2. G8 set: which white cell count supports the diagnosis of appendicitis?',
     ['A. Above 11 ×10⁹/L', 'B. Below 4 ×10⁹/L', 'C. Exactly 7 ×10⁹/L', 'D. Above 30 ×10⁹/L only']),
    ('3. G8 set: which imaging test is preferred in adults when the diagnosis is uncertain?',
     ['A. CT abdomen', 'B. Ultrasound', 'C. Plain abdominal X-ray', 'D. Barium enema']),
    ('4. G8 set: which differential diagnosis is excluded in women of reproductive age?',
     ['A. Ectopic pregnancy', 'B. Gout', 'C. Migraine', 'D. Tinea pedis']),
    ('5. G8 set: how many Alvarado points does right iliac fossa tenderness add?',
     ['A. 1', 'B. 2', 'C. 3', 'D. 0']),
]
# Q3's printed key is WRONG on purpose (the lecture prefers CT in adults) — the owner corrects it (AC-26).
BANK_KEY = '1. A 2. A 3. B 4. A 5. B'

if wanted('g8_key_bank.pdf'):
    c = canvas.Canvas(os.path.join(HERE, 'g8_key_bank.pdf'), pagesize=A4)
    c.setTitle('G8 Revision Bank (TEST FIXTURE)')
    c.setFont('DejaVu', 9)
    c.drawString(60, H - 40, 'G8 Revision Bank - TEST FIXTURE')
    c.drawCentredString(W / 2, 30, '1')
    y = H - 90
    c.setFont('DejaVu-Bold', 15)
    c.drawString(60, y, 'G8 Revision Bank')
    y -= 20
    c.setFont('DejaVu', 9)
    c.drawString(60, y, FIX)
    y -= 26
    for stem, opts in BANK:
        c.setFont('DejaVu', 11)
        c.drawString(60, y, stem)
        y -= 17
        for o in opts:
            c.drawString(60, y, o)
            y -= 17
        y -= 8
    c.setFont('DejaVu-Bold', 13)
    c.drawString(60, y, 'Answers')
    y -= 20
    c.setFont('DejaVu', 11)
    c.drawString(60, y, BANK_KEY)
    c.save()

# ── g8_linked_image.pptx ──
if wanted('g8_linked_image.pptx'):
    tmp = tempfile.mkdtemp(prefix='g8-')
    try:
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(1))
        tb.text_frame.text = 'G8 linked picture - TEST FIXTURE (synthetic; the picture is only a link)'
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        s.shapes.add_picture(os.path.join(GOLDEN, 'flowchart.png'), Inches(1), Inches(1.5), Inches(4), Inches(3))
        base = os.path.join(tmp, 'base.pptx')
        prs.save(base)
        out = os.path.join(HERE, 'g8_linked_image.pptx')
        with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == 'ppt/slides/_rels/slide1.xml.rels':
                    t = data.decode()
                    t = re.sub(r'Target="\.\./media/[^"]+"', f'Target="{TRAP}/g8-pptx-ssrf.png" TargetMode="External"', t)
                    data = t.encode()
                elif item.filename == 'ppt/slides/slide1.xml':
                    data = data.decode().replace('r:embed=', 'r:link=').encode()
                elif item.filename.startswith('ppt/media/'):
                    continue  # the picture exists only as the external link
                zout.writestr(item, data)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

# ── g8_linked_image.doc ──
if wanted('g8_linked_image.doc'):
    tmp = tempfile.mkdtemp(prefix='g8-')
    try:
        html = os.path.join(tmp, 'g8_linked_image.html')
        with open(html, 'w', encoding='utf-8') as f:
            f.write(
                '<html><head><meta charset="utf-8"><title>G8 linked picture (TEST FIXTURE)</title></head><body>'
                f'<p>G8 linked picture - {FIX}</p>'
                f'<p><img src="{TRAP}/g8-doc-ssrf.png" width="120" height="90" alt="linked"></p>'
                '<p>The picture above is only a link to another host.</p></body></html>'
            )
        env = dict(NO_NET, HOME=tmp, TMPDIR=tmp)
        subprocess.run(
            ['soffice', f'-env:UserInstallation=file://{tmp}/profile', '--headless', '--norestore', '--nologo',
             '--convert-to', 'doc:MS Word 97', '--outdir', tmp, html],
            check=True, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180,
        )
        shutil.copyfile(os.path.join(tmp, 'g8_linked_image.doc'), os.path.join(HERE, 'g8_linked_image.doc'))
        with open(os.path.join(HERE, 'g8_linked_image.doc'), 'rb') as f:
            data = f.read()
        if b'65001' not in data and '65001'.encode('utf-16-le') not in data:
            raise SystemExit('the linked URL is not in the .doc (LibreOffice embedded the picture?)')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
print('ok')
